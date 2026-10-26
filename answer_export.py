"""
answer_export.py -- an answer, or a whole conversation, as a PDF or a Word file (built 2026-10-23 for the new front end; owner's go: "I want a PDF / word export").

WHY A NEW MODULE, NOT answer_pdf.py: answer_pdf.py (the WhatsApp "PDF" command) keeps only the answer text, the sections and judgments the text names, and the hedged "other cases". The website page shows more, and every extra box
is a WARNING a reader needs: the currency notes ("reported as overruled, not yet confirmed by our reviewer"), the notes from the checks run by code, the "wider search, not independently verified" boxes, the Act notice, the section's whole
official text. A document outlives the page it came from, so an export that dropped those would show the answer without its warnings. This module therefore takes the SAME data the page draws from (answer_view.build, the `view`), in the page's
own order, and carries every box into the file with its wording. It reuses answer_pdf's bundled fonts, disclaimer and "what's unclear" checklist.

HOW IT WORKS
  1. build_document(entries, ...) -> a neutral document model (a list of blocks: title, headings, paragraphs with bold / italic / link runs, lists, tables, boxes, official-text blocks). One place decides what goes in and in what order.
  2. render_pdf(doc) (reportlab; DejaVu fonts, so rupee signs and curly quotes print) and render_docx(doc) (python-docx) turn the SAME model into the two formats, so the two files cannot differ in content.
  3. verify_pdf / verify_docx read the finished file back and prove every block's text is inside it (letters and digits only, so line breaks and styling cannot fool the check). export() builds, verifies, and for a PDF falls back to a plain
     layout if the check fails. A PDF cannot draw a script its font lacks (Devanagari, for one): export() then refuses with ScriptNotSupported instead of sending a page of empty boxes, and the Word file is offered instead.
  4. Nothing here calls a model or the network; an export costs nothing. The caller (api/routers/chats.py) builds the entries from the SAVED messages of the person's own conversation only, so a document cannot be forged by editing a request.

Nothing in an export says a case is good law. The wording of every box is the page's own (test_answer_export.py checks the boilerplate against recourse_app.py's source, so a change of wording on the page fails the test until this file follows).
"""
import datetime
import html as _html
import io
import logging
import os
import re
import zipfile
from xml.sax.saxutils import escape

import answer_pdf

logger = logging.getLogger("answer_export")

IST = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
MAX_ENTRIES = 12
FORMATS = ("pdf", "docx")
MIME = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}

DISCLAIMER = answer_pdf.DISCLAIMER
DISCLAIMER_LAWYER = ("Recourse's research notes are general legal information, not legal advice, and not a court document. Check every case and every section against the original before you rely on it or cite it.")

# ---- the page's own words (recourse_app.render_answer); test_answer_export.py compares them with the page's source
CONFLICT = "More than one provision applies and they don't all say the same thing — Recourse lays out each rather than picking one for you."
FALLBACK_LEAD = "Here is what the law says on this:"
H_SOURCES = "The sections and judgments this answer rests on"
H_RELIED = "Judgments this answer relies on"
H_CURRENCY = "Currency notes (checked by code, not by the model)"
H_GROUNDING = "Checked by code against the sources we read"
H_SCI_FED = "Supreme Court passages from a wider search that this answer also read (not independently verified)"
H_SCI_SHOWN = "More Supreme Court judgments that may be relevant (a wider search; not independently verified)"
SCI_FOOT = ("These passages came from a search of a much wider judgment database. Nobody at Recourse has read and confirmed them, and this search does not say whether a decision is still good law. "
            "Read the real judgment yourself, or with a lawyer, before relying on it.")
H_ANCHORS = "Verbatim Supreme Court excerpts this answer read"
ANCHORS_UNCONFIRMED = " (not yet confirmed by our reviewer)"
ANCHORS_FOOT_1 = "These are the judgments' own words, checked by a program against the Supreme Court's own PDF. "
ANCHORS_FOOT_2 = "A person at Recourse has not yet confirmed that they are the right authority or that the law has not moved on. "
ANCHORS_FOOT_3 = "Read the real judgment yourself, or with a lawyer, before relying on it."
H_UNVERIFIED = "Other real court cases that might be relevant (not independently verified)"
UNVERIFIED_FOOT = ("These came up in a search of a much wider judgment database. Nobody at Recourse has read and confirmed them yet — read the real judgment yourself, or with a lawyer, "
                   "before relying on it.")
OPEN_BUCKET = "copy from the open Supreme Court judgments bucket (CC-BY-4.0, Dattam Labs)"


class ExportError(Exception):
    """The file could not be built."""


class ScriptNotSupported(ExportError):
    """A PDF cannot draw some character in the text (the bundled font lacks it); the Word file can."""


# ======================================================================================================================
# 1. the document model
# ======================================================================================================================
def R(text, **fmt):
    """One run of text: fmt keys b (bold), i (italic), href (a link), muted (small grey)."""
    d = {"text": "" if text is None else str(text)}
    d.update({k: v for k, v in fmt.items() if v})
    return d


def plain(runs):
    return "".join(r["text"] for r in runs)


def safe_url(url):
    """The address if it is a plain http(s) link, else None (no javascript:, data:, relative or odd addresses)."""
    s = str(url or "").strip()
    if not re.match(r"^https?://", s, re.I) or len(s) > 2000 or re.search(r"[\s<>\"']", s):
        return None
    return s


def link_runs(label, url, **fmt):
    """A link: the label (clickable) and, when it differs from the address, the address in small type -- a printed page cannot be clicked."""
    u = safe_url(url)
    label = "" if label is None else str(label)
    if not u:
        return [R(label, **fmt)] if label else []
    out = [R(label or u, href=u, **fmt)]
    if label and label != u:
        out.append(R(" (" + u + ")", muted=True))
    return out


_INLINE = re.compile(
    r"\*\*(?P<b>(?:(?!\*\*).)+?)\*\*"
    r"|(?<![\w*])\*(?P<i>(?![\s*])(?:(?!\*).)+?)(?<![\s*])\*(?![\w*])"
    r"|\[(?P<lt>[^\]\n]+)\]\((?P<lu>[^)\s]+)\)"
    r"|`(?P<c>[^`\n]+)`", re.S)


def inline(text, fmt=None):
    """Markdown inline marks (**bold**, *italic*, [label](url), `code`) -> runs. Anything else, including a lone asterisk (5*4), stays as written."""
    fmt = dict(fmt or {})
    out, pos = [], 0
    text = "" if text is None else str(text)
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            out.append(R(text[pos:m.start()], **fmt))
        if m.group("b") is not None:
            out += inline(m.group("b"), {**fmt, "b": True})
        elif m.group("i") is not None:
            out += inline(m.group("i"), {**fmt, "i": True})
        elif m.group("lt") is not None:
            out += link_runs(m.group("lt"), m.group("lu"), **fmt)
        else:
            out.append(R(m.group("c"), **fmt))
        pos = m.end()
    if pos < len(text):
        out.append(R(text[pos:], **fmt))
    return [r for r in out if r["text"]]


_BULLET = re.compile(r"^\s{0,3}[-*•]\s+(.*\S)\s*$")
_NUMBERED = re.compile(r"^\s{0,3}(\d{1,3})[.)]\s+(.*\S)\s*$")
_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*\S)\s*$")
_QUOTE = re.compile(r"^\s{0,3}>\s?(.*)$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


def _cells(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.replace("\\|", "|").strip() for c in re.split(r"(?<!\\)\|", line)]


def md_blocks(md):
    """The markdown the page draws -> document blocks: paragraphs, label headings (a paragraph that is only a bold line), bullet and numbered lists, quotes, pipe tables. Soft line breaks join with a space (as the page does)."""
    blocks, para, lines = [], [], (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")

    def flush():
        if not para:
            return
        text = " ".join(s.strip() for s in para).strip()
        para.clear()
        if not text:
            return
        runs = inline(text)
        if len(runs) == 1 and runs[0].get("b") and not runs[0].get("href"):
            blocks.append({"t": "h3", "text": runs[0]["text"].strip().rstrip(":").strip()})
        else:
            blocks.append({"t": "p", "runs": runs})

    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            flush()
            i += 1
            continue
        m = _HEADING.match(ln)
        if m:
            flush()
            blocks.append({"t": "h3", "text": plain(inline(m.group(1)))})
            i += 1
            continue
        if ln.strip().startswith("|") and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            flush()
            head = [plain(inline(c)) for c in _cells(ln)]
            rows, i = [], i + 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                cs = _cells(lines[i])
                cs = (cs + [""] * len(head))[:len(head)]
                rows.append([inline(c) for c in cs])
                i += 1
            blocks.append({"t": "table", "head": head, "rows": rows})
            continue
        m = _BULLET.match(ln)
        n = _NUMBERED.match(ln)
        if m or n:
            flush()
            ordered = bool(n)
            items = []
            while i < len(lines):
                a, b = _BULLET.match(lines[i]), _NUMBERED.match(lines[i])
                if ordered and b:
                    items.append([int(b.group(1)), b.group(2)])
                elif not ordered and a:
                    items.append([None, a.group(1)])
                elif lines[i].strip() and lines[i].startswith("  ") and items and not (a or b):
                    items[-1][1] += " " + lines[i].strip()
                else:
                    break
                i += 1
            blocks.append({"t": "ol" if ordered else "ul", "items": [{"num": num, "runs": inline(txt)} for num, txt in items]})
            continue
        m = _QUOTE.match(ln)
        if m:
            flush()
            q = []
            while i < len(lines) and _QUOTE.match(lines[i]):
                q.append(_QUOTE.match(lines[i]).group(1).strip())
                i += 1
            blocks.append({"t": "quote", "runs": inline(" ".join(x for x in q if x))})
            continue
        para.append(ln)
        i += 1
    flush()
    return blocks


def _p(text, **fmt):
    return {"t": "p", "runs": [R(text, **fmt)]}


def _box(runs, label=None):
    return {"t": "box", "label": label, "runs": runs}


def _official(text):
    t = re.sub(r"\n{3,}", "\n\n", str(text or "")).strip("\n")
    return {"t": "mono", "text": t}


def html_to_runs(html_text):
    """The Act notice is written by the engine as escaped HTML with a few tags (b, a). -> runs: the text, bold kept, a link's address printed after its words."""
    src = str(html_text or "")
    out, pos = [], 0
    for m in re.finditer(r"<a\s[^>]*?href=[\"']([^\"']*)[\"'][^>]*>(.*?)</a>|<b>(.*?)</b>|<strong>(.*?)</strong>|<[^>]+>", src, re.I | re.S):
        if m.start() > pos:
            out.append(R(_html.unescape(src[pos:m.start()])))
        if m.group(0).lower().startswith("<a") and m.group(1) is not None:
            label = _html.unescape(re.sub(r"<[^>]+>", "", m.group(2) or ""))
            out += link_runs(label, _html.unescape(m.group(1)))
        elif m.group(3) is not None or m.group(4) is not None:
            out.append(R(_html.unescape(re.sub(r"<[^>]+>", "", m.group(3) or m.group(4) or "")), b=True))
        pos = m.end()
    if pos < len(src):
        out.append(R(_html.unescape(src[pos:])))
    runs = [r for r in out if r["text"]]
    if runs:
        runs[0]["text"] = runs[0]["text"].lstrip()
        runs[-1]["text"] = runs[-1]["text"].rstrip()
    return [r for r in runs if r["text"]]


def _fmt_time(dt):
    if isinstance(dt, (int, float)):
        dt = datetime.datetime.fromtimestamp(dt, tz=IST)
    if dt is None:
        return ""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=datetime.timezone.utc)
    return dt.astimezone(IST).strftime("%d %b %Y, %H:%M IST")


def _s(v):
    return None if v is None else str(v)


# ---- the blocks for one answer, in the page's own order ----------------------------------------------------------------
def _answer_blocks(view, lawyer):
    B = []
    if view.get("conflicting"):
        B.append(_p(CONFLICT, i=True, muted=True))
    if view.get("outside_scope_note"):
        B.append(_box([R(view["outside_scope_note"])]))
    if view.get("answer_markdown"):
        B += md_blocks(view["answer_markdown"])
    else:
        B += [_p(FALLBACK_LEAD), {"t": "quote", "runs": [R(view.get("fallback_text") or "")]}]
    if not lawyer:
        try:
            cl = answer_pdf.unclear_checklist(view.get("copy_text") or "") if view.get("answer_markdown") else {"items": [], "documents": [], "from_answer": ""}
        except Exception:  # noqa: BLE001
            cl = {"items": [], "documents": [], "from_answer": ""}
        if cl["items"] or cl["from_answer"] or cl["documents"]:
            B.append({"t": "h3", "text": "Facts to find out or bring to a lawyer"})
            B.append(_p("Taken from what the answer itself said was unclear.", muted=True))
            for it in cl["items"]:
                B.append(_p("☐  " + it))
            if cl["from_answer"]:
                B.append(_p("From the answer: " + cl["from_answer"]))
            if cl["documents"]:
                B.append({"t": "h3", "text": "Documents to have ready"})
                for d in cl["documents"]:
                    B.append(_p("☐  " + d))
    if view.get("act_notice_html"):
        runs = html_to_runs(view["act_notice_html"])
        if runs:
            B.append(_box(runs))
    for it in view.get("section_notes") or []:
        if it.get("lead"):
            B.append(_box([R(it["lead"])]))
        for t in it.get("texts") or []:
            B.append(_official(t))
        for c in it.get("cites") or []:
            if c.get("lead"):
                B.append(_box([R(c["lead"])]))
            for t in c.get("texts") or []:
                B.append(_official(t))
        cap = [R(it.get("caption") or "", muted=True)]
        if safe_url(it.get("url")):
            cap += [R(" · ", muted=True)] + link_runs("India Code", it["url"], muted=True)
        B.append({"t": "p", "runs": cap})
    auth = view.get("authorities")
    if auth:
        for t in auth.get("tables") or []:
            B.append({"t": "h3", "text": t.get("heading") or ""})
            B += [b for b in md_blocks(t.get("markdown") or "") if b["t"] == "table"] or md_blocks(t.get("markdown") or "")
        if auth.get("note"):
            B.append(_p(auth["note"], muted=True))
    if view.get("statutes") or view.get("judgments"):
        B.append({"t": "h2", "text": H_SOURCES})
        for s in view.get("statutes") or []:
            B.append({"t": "h3", "text": s.get("heading") or ""})
            if s.get("caption"):
                cap = [R(s["caption"], muted=True)]
                if safe_url(s.get("caption_url")):
                    cap += [R(" · ", muted=True)] + link_runs("India Code", s["caption_url"], muted=True)
                B.append({"t": "p", "runs": cap})
            B.append(_official(s.get("text") or ""))
        if view.get("judgments"):
            B.append({"t": "h3", "text": "Judgments"})
        for c in view.get("judgments") or []:
            B.append({"t": "p", "runs": [R(c.get("name") or "", b=True)] + ([R("  ·  " + c["citation"])] if c.get("citation") else [])})
            for cv in c.get("caveats") or []:
                B.append(_box([R(cv)]))
            for e in c.get("excerpts") or []:
                B.append(_p(("paragraph " + e["paragraph"]) if e.get("paragraph") else "excerpt", i=True, muted=True))
                B.append(_official(e.get("text") or ""))
    if view.get("cited_links"):
        B.append({"t": "h2", "text": H_RELIED})
        B.append({"t": "ul", "items": [{"num": None, "runs": link_runs(c.get("case_name") or "", c.get("url")) + [R("  ·  " + (c.get("source_label") or "source link"), muted=True)]} for c in view["cited_links"]]})
    reg = view.get("register")
    if reg:
        B.append({"t": "h2", "text": H_CURRENCY})
        for n in list(reg.get("notes") or []) + list(reg.get("possible") or []):
            B.append(_box([R(n)]))
        if reg.get("not_checked"):
            B.append(_p(reg["not_checked"], muted=True))
    if view.get("grounding_notes"):
        B.append({"t": "h2", "text": H_GROUNDING})
        for n in view["grounding_notes"]:
            B.append(_box([R(n)]))
    sci = view.get("sci_related")
    if sci and sci.get("items"):
        import sci_search_tier as _SCI
        tagged = bool(sci.get("court_tags"))                      # a court is named beside a passage only when some passage of the list is not the Supreme Court's
        scope = sci.get("scope") or "Supreme Court"
        B.append({"t": "h2", "text": (H_SCI_FED if sci.get("fed") else H_SCI_SHOWN) if scope == "Supreme Court" else _SCI.box_title(None, bool(sci.get("fed")), scope)})
        B.append(_p(SCI_FOOT + (_SCI.COURT_FOOT if tagged else "") + _SCI.asof_sentence(sci.get("as_of"), sci.get("as_of_by_court")), muted=True))
        for j in sci["items"]:
            head = link_runs(j.get("case_name") or "", j.get("link"), b=True)
            if not head:
                head = [R(j.get("case_name") or "", b=True)]
            credit = _SCI.credit_for(j.get("link_kind"))
            tail = ((" (" + j["date"] + ")") if j.get("date") else "") + ((" · " + j["court"]) if tagged and j.get("court") else "") + ((" — paragraph " + j["paragraph"]) if j.get("paragraph") else "") + ((" · " + credit) if credit else "")
            B.append({"t": "p", "runs": head + ([R(tail)] if tail else [])})
            if j.get("register_note"):
                B.append(_box([R(j["register_note"])]))
            if j.get("passage_note"):
                B.append(_p(j["passage_note"], muted=True))
            if j.get("names_register_note"):
                B.append(_box([R("This passage names a case on our register: " + j["names_register_note"])]))
            if j.get("excerpt"):
                B.append(_official(j["excerpt"]))
    an = view.get("anchors")
    if an and an.get("items"):
        B.append({"t": "h2", "text": H_ANCHORS + (ANCHORS_UNCONFIRMED if an.get("unconfirmed") else "")})
        B.append(_p(ANCHORS_FOOT_1 + (ANCHORS_FOOT_2 if an.get("unconfirmed") else "") + ANCHORS_FOOT_3, muted=True))
        for a in an["items"]:
            head = link_runs(a.get("case_name") or "", a.get("link"), b=True) or [R(a.get("case_name") or "", b=True)]
            tail = ((" (" + a["year"] + ")") if a.get("year") else "") + ((", " + a["citation"]) if a.get("citation") else "")
            B.append({"t": "p", "runs": head + ([R(tail)] if tail else [])})
            for e in a.get("excerpts") or []:
                B.append(_p((e.get("label") or "") + ((" — paragraph " + e["paragraph"]) if e.get("paragraph") else ""), i=True, muted=True))
                B.append(_official(e.get("text") or ""))
                if e.get("note"):
                    B.append(_p(e["note"], muted=True))
    if view.get("unverified"):
        B.append({"t": "h2", "text": H_UNVERIFIED})
        B.append(_p(UNVERIFIED_FOOT, muted=True))
        for j in view["unverified"]:
            runs = link_runs(j.get("case_name") or "", j.get("url")) or [R(j.get("case_name") or "")]
            if j.get("paragraph"):
                runs.append(R(" — paragraph " + j["paragraph"]))
            if j.get("procedural"):
                runs.append(R(" — may be a procedural/bail order, not a full judgment"))
            B.append({"t": "p", "runs": runs})
    return B


def _notice_blocks(view):
    n = view.get("notice") or {}
    B = []
    if n.get("label"):
        B.append(_p(n["label"], b=True, muted=True))
    if n.get("title"):
        B.append({"t": "h3", "text": n["title"]})
    if n.get("state") == "covered_elsewhere_in_tool" and n.get("body_template"):
        B.append(_box([R(n["body_template"].replace("{topic}", str(n.get("topic") or "")))]))
    elif n.get("state") == "adjacent_uncovered":
        B.append(_box([R((n.get("body") or "") + ((" " + n["reasoning"]) if n.get("reasoning") else ""))]))
        if n.get("footer"):
            B.append(_p(n["footer"], muted=True))
    else:
        B.append(_box([R(n.get("body") or "")]))
    if view.get("act_notice_html"):
        runs = html_to_runs(view["act_notice_html"])
        if runs:
            B.append(_box(runs))
    notes = view.get("register_notes") or []
    if notes:
        B.append({"t": "h2", "text": H_CURRENCY})
        for x in notes:
            B.append(_box([R(x)]))
    return B


_CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufffe\uffff]")


def _clean(x):
    """Every string in the model without the control characters that a judgment's scanned text can carry (the Word library refuses them, the PDF library draws boxes); form feeds become line breaks."""
    if isinstance(x, str):
        return _CONTROL.sub("", x.replace("\x0c", "\n"))
    if isinstance(x, list):
        return [_clean(v) for v in x]
    if isinstance(x, dict):
        return {k: _clean(v) for k, v in x.items()}
    return x


def build_document(entries, *, kind="answer", title_chat=None, generated_at=None):
    """entries: [{'question': str, 'view': dict|None (answer_view.build), 'plain': str|None (answer text when no view), 'answer_id': str, 'created_at': datetime|float|None, 'audience': 'layman'|'lawyer'}], oldest first.
    -> {'title', 'blocks', 'footer', 'meta', 'lawyer', 'count'}."""
    entries = list(entries or [])[-MAX_ENTRIES:]
    if not entries:
        raise ExportError("nothing to export")
    lawyer = any(e.get("audience") == "lawyer" for e in entries)
    multi = len(entries) > 1
    generated = _fmt_time(generated_at if generated_at is not None else datetime.datetime.now(datetime.timezone.utc))
    if multi or kind == "conversation":
        title = "Your conversation with Recourse" if not lawyer else "Research notes from Recourse"
        ids = f"{len(entries)} answers"
        meta = f"Recourse · {len(entries)} questions and answers · generated {generated}"
    else:
        title = ("Your question and the answer" if not lawyer else "Research note")
        ids = "Answer ID " + entries[0]["answer_id"]
        meta = f"Recourse · answer ID {entries[0]['answer_id']} · generated {generated}"
    B = [{"t": "title", "text": title}, {"t": "meta", "text": meta}]
    if title_chat:
        B.append(_p("Conversation: " + title_chat, muted=True))
    B.append({"t": "note", "runs": [R(DISCLAIMER_LAWYER if lawyer else DISCLAIMER, i=True)]})
    B.append({"t": "rule"})
    if multi:
        B.append({"t": "h2", "text": "Contents"})
        for n, e in enumerate(entries, 1):
            q = (e.get("question") or "").strip()
            q = q if len(q) <= 95 else q[:92].rstrip() + "..."
            B.append(_p(f"Question {n} — {q}  ({e['answer_id']})", muted=True))
    for n, e in enumerate(entries, 1):
        B.append({"t": "h2", "text": f"Question {n}" if multi else "Your question"})
        if multi:
            B.append(_p(f"Answer ID {e['answer_id']} · {_fmt_time(e.get('created_at'))}", muted=True))
        B.append({"t": "quote", "runs": [R((e.get("question") or "").strip() or "(no text)")]})
        B.append({"t": "h2" if not multi else "h3", "text": "The answer" if not multi else "Answer"})
        v = e.get("view")
        if isinstance(v, dict) and v.get("kind") == "answer":
            B += _answer_blocks(v, e.get("audience") == "lawyer")
        elif isinstance(v, dict) and v.get("kind") == "notice":
            B += _notice_blocks(v)
        elif e.get("plain"):
            B += md_blocks(e["plain"])
        else:
            B.append(_box([R("No answer was saved for this question.")]))
    if not lawyer:
        B.append({"t": "h2", "text": "Notes for my lawyer"})
        B.append({"t": "lines", "n": 4})
    return {"title": _clean(title), "blocks": _clean(B), "footer": _clean(f"Recourse — information, not legal advice · {ids}"), "meta": _clean(meta), "lawyer": lawyer, "count": len(entries)}


# ---- what a reader of the finished file must be able to find ------------------------------------------------------------
def expected_texts(doc):
    out = []
    for b in doc["blocks"]:
        t = b["t"]
        if t in ("title", "meta", "h2", "h3", "mono"):
            out.append(b["text"])
        elif t in ("p", "note", "quote", "box"):
            out.append(plain(b["runs"]))
        elif t in ("ul", "ol"):
            out += [plain(it["runs"]) for it in b["items"]]
        elif t == "table":
            out += list(b["head"])
            for row in b["rows"]:
                out += [plain(c) for c in row]
    return [x for x in out if _norm(x)]


def _norm(s):
    return "".join(ch.lower() for ch in str(s or "") if ch.isalnum())


def all_text(doc):
    return "\n".join(expected_texts(doc))


# ======================================================================================================================
# 2. the PDF
# ======================================================================================================================
_FOOTER_CLIP_PT = 62


def _font_covers():
    """A function char -> bool: does the bundled DejaVu font have a glyph for it? None when the font cannot be read."""
    try:
        from reportlab.pdfbase.ttfonts import TTFont
        face = TTFont("DejaVuProbe", os.path.join(answer_pdf._FONT_DIR, "DejaVuSans.ttf")).face
        cmap = face.charToGlyph
        return lambda ch: ord(ch) in cmap
    except Exception:  # noqa: BLE001
        return None


def uncovered_characters(doc):
    """The characters in the text that the PDF font cannot draw (letters of other scripts, mostly). Spaces and control characters are ignored."""
    covers = _font_covers()
    if covers is None:
        return []
    bad = []
    for ch in sorted({c for t in expected_texts(doc) + [doc["footer"]] for c in t if not c.isspace() and ord(c) >= 32}):
        if not covers(ch):
            bad.append(ch)
    return bad


def _rl_runs(runs):
    out = []
    for r in runs:
        t = escape(r["text"]).replace("\n", "<br/>")
        if r.get("b"):
            t = f"<b>{t}</b>"
        if r.get("i"):
            t = f"<i>{t}</i>"
        if r.get("muted"):
            t = f'<font color="#5B6675" size="8">{t}</font>'
        if r.get("href"):
            t = f'<link href="{escape(r["href"], {chr(34): "&quot;"})}" color="#1F4E8C">{t}</link>'
        out.append(t)
    return "".join(out)


def render_pdf(doc, plain_layout=False):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fonts_ok = answer_pdf._register_fonts()
    base, bold, ital = ("DejaVu", "DejaVu-Bold", "DejaVu-Italic") if fonts_ok else ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")
    INK, MUTED, ACCENT, PALE, WARM = colors.HexColor("#1B2430"), colors.HexColor("#5B6675"), colors.HexColor("#1F4E8C"), colors.HexColor("#F3F5F8"), colors.HexColor("#FFF6E0")

    def S(name, **kw):
        d = dict(fontName=base, fontSize=10, leading=14.5, textColor=INK, alignment=TA_LEFT, spaceAfter=5)
        d.update(kw)
        return ParagraphStyle(name, **d)

    st = {
        "title": S("title", fontName=bold, fontSize=17, leading=22, spaceAfter=3),
        "meta": S("meta", fontSize=8.5, leading=12, textColor=MUTED, spaceAfter=8),
        "h2": S("h2", fontName=bold, fontSize=12, leading=16, textColor=ACCENT, spaceBefore=12, spaceAfter=4, keepWithNext=1),
        "h3": S("h3", fontName=bold, fontSize=10.5, leading=14, spaceBefore=6, spaceAfter=2, keepWithNext=1),
        "body": S("body"),
        "item": S("item", leftIndent=16, bulletIndent=2, spaceAfter=3),
        "quote": S("quote", fontSize=10.5, leading=15, backColor=PALE, borderPadding=(6, 6, 6, 6), spaceBefore=4, spaceAfter=10),
        "note": S("note", fontName=ital, textColor=MUTED, fontSize=9.5),
        "box": S("box", fontSize=9.5, leading=13.5, backColor=WARM, borderPadding=(5, 5, 5, 5), spaceBefore=3, spaceAfter=9),
        "mono": S("mono", fontSize=8.5, leading=12, backColor=PALE, borderPadding=(5, 5, 5, 5), spaceAfter=9),
        "cell": S("cell", fontSize=7.8, leading=10.2, spaceAfter=0),
        "cellh": S("cellh", fontName=bold, fontSize=7.8, leading=10.2, spaceAfter=0),
        "lines": S("lines", fontSize=10, leading=22),
    }

    def para(runs, style, bullet=None):
        if plain_layout:
            runs = [{"text": r["text"]} for r in runs]
        return Paragraph(_rl_runs(runs) or "&nbsp;", st[style], bulletText=bullet)

    story = []
    for b in doc["blocks"]:
        t = b["t"]
        if t == "title":
            story.append(Paragraph(escape(b["text"]), st["title"]))
        elif t == "meta":
            story.append(Paragraph(escape(b["text"]), st["meta"]))
        elif t == "rule":
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#C9D0DA"), spaceBefore=6, spaceAfter=6))
        elif t in ("h2", "h3"):
            story.append(Paragraph(escape(b["text"]), st[t]))
        elif t == "p":
            story.append(para(b["runs"], "body"))
        elif t == "note":
            story.append(para(b["runs"], "note"))
        elif t == "quote":
            story.append(para(b["runs"], "quote"))
        elif t == "box":
            runs = ([R(b["label"] + ": ", b=True)] if b.get("label") else []) + b["runs"]
            story.append(para(runs, "box"))
        elif t == "mono":
            story.append(Paragraph(escape(b["text"]).replace("\n", "<br/>"), st["mono"]))
        elif t == "ul":
            for it in b["items"]:
                story.append(para(it["runs"], "item", bullet="•"))
        elif t == "ol":
            for k, it in enumerate(b["items"], 1):
                story.append(para(it["runs"], "item", bullet=f"{it['num'] if it.get('num') is not None else k}."))
        elif t == "lines":
            for _ in range(b.get("n", 4)):
                story.append(HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#C9D0DA"), spaceBefore=14, spaceAfter=0))
        elif t == "table":
            if plain_layout:
                story.append(Paragraph(escape(" | ".join(b["head"])), st["body"]))
                for row in b["rows"]:
                    story.append(Paragraph(escape(" | ".join(plain(c) for c in row)), st["body"]))
                continue
            width = A4[0] - 36 * mm
            weights = []
            for h in b["head"]:
                hl = h.lower()
                weights.append(26 if hl == "case" else 15 if hl == "citation" else 9 if hl in ("paragraph", "bench") else 32 if "stands" in hl else 18 if hl == "link" else 100 / max(1, len(b["head"])))
            tot = sum(weights)
            cols = [width * w / tot for w in weights]
            data = [[Paragraph(escape(h), st["cellh"]) for h in b["head"]]]
            for row in b["rows"]:
                data.append([para(c, "cell") for c in row])
            tbl = Table(data, colWidths=cols, repeatRows=1, hAlign="LEFT")
            tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), PALE), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D0DA")),
                                     ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
            story += [tbl, Spacer(1, 8)]
    footer = doc["footer"]

    class _Numbered(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self.setFont(base, 7.5)
                self.setFillColor(MUTED)
                self.drawString(18 * mm, 13 * mm, footer)
                self.drawRightString(A4[0] - 18 * mm, 13 * mm, f"Page {self._pageNumber} of {n}")
                super().showPage()
            super().save()

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=26 * mm, title="Recourse — " + doc["title"].lower(), author="Recourse", subject=doc["meta"])
    pdf.build(story, canvasmaker=_Numbered)
    return buf.getvalue()


def verify_pdf(data, doc):
    """{'ok': True/False/None, 'checked': n, 'missing': [...]}: every block's text is in the PDF (the footer is excluded, so a block that crosses a page is still found). None when no PDF reader is available."""
    try:
        import pymupdf as fitz
    except Exception:  # noqa: BLE001
        try:
            import fitz
        except Exception:  # noqa: BLE001
            return {"ok": None, "checked": 0, "missing": [], "reason": "no PDF reader available"}
    try:
        d = fitz.open(stream=data, filetype="pdf")
        parts = []
        for page in d:
            r = page.rect
            parts.append(page.get_text("text", clip=fitz.Rect(0, 0, r.width, r.height - _FOOTER_CLIP_PT)))
        hay = _norm("\n".join(parts))
        exp = expected_texts(doc)
        missing = [e for e in exp if _norm(e) not in hay]
        return {"ok": not missing, "checked": len(exp), "missing": missing}
    except Exception as exc:  # noqa: BLE001
        return {"ok": None, "checked": 0, "missing": [], "reason": f"{type(exc).__name__}: {exc}"}


# ======================================================================================================================
# 3. the Word file
# ======================================================================================================================
def _docx_modules():
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    return docx, WD_ALIGN_PARAGRAPH, RT, OxmlElement, qn, Mm, Pt, RGBColor


_PPR_ORDER = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
              "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
              "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def render_docx(doc):
    docx, WD_ALIGN, RT, OxmlElement, qn, Mm, Pt, RGBColor = _docx_modules()
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(22)
    INK, MUTED, ACCENT = RGBColor(0x1B, 0x24, 0x30), RGBColor(0x5B, 0x66, 0x75), RGBColor(0x1F, 0x4E, 0x8C)

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        s = d.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = ACCENT if name != "Title" else INK
        rpr = s.element.get_or_add_rPr()
        if rpr.rFonts is None:
            rpr._add_rFonts()
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rpr.rFonts.set(qn(a), "Calibri")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rpr.rFonts.get(qn(a)) is not None:
                del rpr.rFonts.attrib[qn(a)]

    def ppr_insert(p, el, tag):
        """Put a paragraph-property element where the file format wants it (Word is strict about the order of these children)."""
        pPr = p._p.get_or_add_pPr()
        pPr.insert_element_before(el, *_PPR_ORDER[_PPR_ORDER.index(tag) + 1:])

    def shade(p, fill, left_color=None):
        if left_color:
            bdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "6"), ("w:color", left_color)):
                left.set(qn(k), v)
            bdr.append(left)
            ppr_insert(p, bdr, "w:pBdr")
        shd = OxmlElement("w:shd")
        for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
            shd.set(qn(k), v)
        ppr_insert(p, shd, "w:shd")

    def bottom_rule(p, between=False):
        bdr = OxmlElement("w:pBdr")
        for tag in (("w:bottom", "w:between") if between else ("w:bottom",)):
            b = OxmlElement(tag)
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "C9D0DA")):
                b.set(qn(k), v)
            bdr.append(b)
        ppr_insert(p, bdr, "w:pBdr")

    def add_run(p, r, size=None, base_bold=False):
        text = r["text"]
        if r.get("href"):
            rid = p.part.relate_to(r["href"], RT.HYPERLINK, is_external=True)
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), rid)
            run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            if r.get("b") or base_bold:
                rpr.append(OxmlElement("w:b"))
            if r.get("i"):
                rpr.append(OxmlElement("w:i"))
            col = OxmlElement("w:color")
            col.set(qn("w:val"), "1F4E8C")
            rpr.append(col)
            if size:
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), str(int(size * 2)))
                rpr.append(sz)
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rpr.append(u)
            run.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            run.append(t)
            h.append(run)
            p._p.append(h)
            return
        parts = text.split("\n")
        for k, part in enumerate(parts):
            run = p.add_run(part)
            run.bold = bool(r.get("b") or base_bold) or None
            run.italic = bool(r.get("i")) or None
            if r.get("muted"):
                run.font.color.rgb = MUTED
                run.font.size = Pt(8.5 if not size else size)
            elif size:
                run.font.size = Pt(size)
            if k < len(parts) - 1:
                run.add_break()

    def para(runs, style=None, size=None, space_after=None, indent=None, hanging=None, italic_all=False):
        p = d.add_paragraph(style=style)
        for r in runs:
            add_run(p, {**r, "i": r.get("i") or italic_all}, size=size)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if indent is not None:
            p.paragraph_format.left_indent = Mm(indent)
        if hanging is not None:
            p.paragraph_format.first_line_indent = Mm(-hanging)
        return p

    for b in doc["blocks"]:
        t = b["t"]
        if t == "title":
            p = d.add_paragraph(b["text"], style="Title")
            p.paragraph_format.space_after = Pt(3)
        elif t == "meta":
            para([R(b["text"], muted=True)], space_after=8)
        elif t == "rule":
            p = d.add_paragraph()
            bottom_rule(p)
            p.paragraph_format.space_after = Pt(6)
        elif t == "h2":
            p = d.add_paragraph(b["text"], style="Heading 1")
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
        elif t == "h3":
            p = d.add_paragraph(b["text"], style="Heading 2")
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
        elif t == "p":
            para(b["runs"])
        elif t == "note":
            para(b["runs"], italic_all=True)
        elif t == "quote":
            p = para(b["runs"], space_after=10)
            shade(p, "F3F5F8", "1F4E8C")
        elif t == "box":
            runs = ([R(b["label"] + ": ", b=True)] if b.get("label") else []) + b["runs"]
            p = para(runs, size=9.5, space_after=8)
            shade(p, "FFF6E0", "D9A400")
        elif t == "mono":
            p = d.add_paragraph()
            add_run(p, R(b["text"]), size=9)
            p.paragraph_format.space_after = Pt(9)
            shade(p, "F3F5F8")
        elif t == "ul":
            for it in b["items"]:
                para(it["runs"], style="List Bullet", space_after=3)
        elif t == "ol":
            for k, it in enumerate(b["items"], 1):
                num = it["num"] if it.get("num") is not None else k
                para([R(f"{num}.  ")] + it["runs"], space_after=3, indent=7, hanging=5)
        elif t == "lines":
            for _ in range(b.get("n", 4)):
                p = d.add_paragraph()
                bottom_rule(p, between=True)
                p.paragraph_format.space_after = Pt(10)
        elif t == "table":
            ncols = len(b["head"])
            tbl = d.add_table(rows=1, cols=ncols)
            tbl.style = "Table Grid"
            tbl.autofit = False
            total_mm = 170
            weights = []
            for h in b["head"]:
                hl = h.lower()
                weights.append(26 if hl == "case" else 15 if hl == "citation" else 9 if hl in ("paragraph", "bench") else 32 if "stands" in hl else 18 if hl == "link" else 100 / max(1, ncols))
            tw = sum(weights)
            widths = [Mm(total_mm * w / tw) for w in weights]

            def fill_cell(cell, runs, header=False, width=None):
                if width is not None:
                    cell.width = width
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                for r in runs:
                    add_run(p, r, size=8, base_bold=header)
                if header:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "F3F5F8")):
                        shd.set(qn(k), v)
                    tcPr.append(shd)

            for k, h in enumerate(b["head"]):
                fill_cell(tbl.rows[0].cells[k], [R(h)], header=True, width=widths[k])
            for row in b["rows"]:
                cells = tbl.add_row().cells
                for k, c in enumerate(row):
                    fill_cell(cells[k], c, width=widths[k])
            hdr = tbl.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            hdr.append(th)
            d.add_paragraph().paragraph_format.space_after = Pt(4)

    # footer with the page number
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(doc["footer"] + " · Page ")
    fr.font.size = Pt(7.5)
    fr.font.color.rgb = MUTED

    def field(par, instr):
        """A Word field (PAGE, NUMPAGES): begin, instruction, separate, a placeholder result, end, each in its own run so Word and LibreOffice both read it."""
        def piece():
            run = par.add_run()
            run.font.size = Pt(7.5)
            run.font.color.rgb = MUTED
            return run

        r1 = piece()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "begin")
        r1._r.append(fc)
        r2 = piece()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = " " + instr + " "
        r2._r.append(it)
        r3 = piece()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "separate")
        r3._r.append(fc)
        r4 = piece()
        r4.text = "1"
        r5 = piece()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "end")
        r5._r.append(fc)

    field(fp, "PAGE")
    fr2 = fp.add_run(" of ")
    fr2.font.size = Pt(7.5)
    fr2.font.color.rgb = MUTED
    field(fp, "NUMPAGES")

    cp = d.core_properties
    cp.title = "Recourse — " + doc["title"].lower()
    cp.author = "Recourse"
    cp.subject = doc["meta"]
    cp.created = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def verify_docx(data, doc):
    """The same check for the Word file: every block's text is found when the file is opened again."""
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    parts += [p.text for p in c.paragraphs]
        # a hyperlink's words live inside <w:hyperlink>, which python-docx's paragraph.text skips in older versions: read the XML text as well
        xml_text = []
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            body = z.read("word/document.xml").decode("utf-8", "replace")
        for m in re.finditer(r"<w:t[^>]*>([^<]*)</w:t>", body):
            xml_text.append(_html.unescape(m.group(1)))
        hay = _norm("\n".join(parts) + "\n" + "".join(xml_text))
        exp = expected_texts(doc)
        missing = [e for e in exp if _norm(e) not in hay]
        return {"ok": not missing, "checked": len(exp), "missing": missing}
    except Exception as exc:  # noqa: BLE001
        return {"ok": None, "checked": 0, "missing": [], "reason": f"{type(exc).__name__}: {exc}"}


# ======================================================================================================================
# 4. the safe build
# ======================================================================================================================
def export(doc, fmt):
    """(bytes, info) for fmt 'pdf' or 'docx'. A PDF is built, proved complete and, if the proof fails, rebuilt in the plain layout. Raises ScriptNotSupported for a PDF whose text has characters the font lacks,
    ExportError if nothing could be built. info = {'ok': True/None, 'fallback': bool, 'checked': n}; 'ok' False is never returned: a file that fails its proof twice is an ExportError."""
    if fmt not in FORMATS:
        raise ExportError("unknown format")
    if fmt == "docx":
        try:
            data = render_docx(doc)
        except Exception as exc:  # noqa: BLE001
            logger.exception("docx export failed")
            raise ExportError(type(exc).__name__) from None
        v = verify_docx(data, doc)
        if v["ok"] is False:
            logger.error("docx export failed its completeness check (%s missing)", len(v["missing"]))
            raise ExportError("incomplete")
        return data, {"ok": v["ok"], "fallback": False, "checked": v["checked"]}
    bad = uncovered_characters(doc)
    if bad:
        raise ScriptNotSupported("".join(bad[:20]))
    first_err = None
    try:
        data = render_pdf(doc)
        v = verify_pdf(data, doc)
        if v["ok"] is not False:
            return data, {"ok": v["ok"], "fallback": False, "checked": v["checked"]}
        logger.warning("styled PDF failed its completeness check (%s missing); rebuilding plain", len(v["missing"]))
    except Exception as exc:  # noqa: BLE001
        logger.exception("styled PDF failed")
        first_err = type(exc).__name__
    try:
        data = render_pdf(doc, plain_layout=True)
        v = verify_pdf(data, doc)
        if v["ok"] is False:
            logger.error("plain PDF ALSO failed its completeness check (%s missing)", len(v["missing"]))
            raise ExportError("incomplete")
        return data, {"ok": v["ok"], "fallback": True, "checked": v["checked"]}
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.exception("plain PDF failed")
        raise ExportError(first_err or type(exc).__name__) from None
