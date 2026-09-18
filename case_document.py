"""
case_document.py -- turns a fetched judgment (case_lookup.fetch_case_text)
into a clean PDF or Word file for WhatsApp delivery.

The disclaimer is printed INSIDE the file, not only in the chat caption:
a PDF gets saved, forwarded and printed separately from the conversation
that produced it, so the warning has to travel with it. Nothing is ever
invented to fill a gap -- a missing citation or source link is printed as
"not available in this record" (the same honest-gap rule used for the
curated anchors, e.g. Prabha Tyagi's missing citation).

reportlab's built-in fonts only cover Latin-1, so typographic characters
(curly quotes, dashes, the rupee sign) are mapped to close equivalents and
anything else unencodable is replaced -- judgments in Indian languages are
NOT rendered faithfully by this PDF and the caller should say so.
"""
import io
import re

DISCLAIMER = (
    "This is a copy of the judgment as retrieved from an open dataset (Vaquill, CC BY 4.0). "
    "It has NOT been independently verified against the official law reporter, and the text may "
    "contain extraction errors. Confirm the citation and the wording against the official source "
    "before relying on this in any filing or proceeding."
)

_SUBST = {
    "‘": "'", "’": "'", "‚": ",", "“": '"', "”": '"',
    "–": "-", "—": "--", "―": "--", "…": "...", " ": " ",
    "•": "-", "₹": "Rs. ", "−": "-", "​": "", "﻿": "",
}


def _safe(text: str) -> str:
    for a, b in _SUBST.items():
        text = text.replace(a, b)
    return text.encode("latin-1", "replace").decode("latin-1")


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _paragraphs(text: str) -> list:
    """Blocks separated by blank lines are paragraphs; single line breaks
    inside a block are extraction wrapping, so they collapse to spaces."""
    blocks = re.split(r"\n\s*\n", text or "")
    return [re.sub(r"\s+", " ", b).strip() for b in blocks if b.strip()]


def _meta_lines(case: dict, doc: dict) -> list:
    from case_lookup import clean_title
    lines = [
        ("Case", clean_title(case.get("title") or "")),
        ("Court", case.get("court") or "not available in this record"),
        ("Decided", case.get("decision_date") or "not available in this record"),
        ("Citation", case.get("citation") or "not available in this record"),
        ("Source link", case.get("source_url") or "not available in this record"),
    ]
    return lines


def _notices(doc: dict) -> list:
    notes = []
    if doc.get("is_procedural"):
        notes.append(
            "NOTE: this record looks like an interim or procedural order (for example a bail or "
            "adjournment order), not a final judgment. Check you have the case you meant."
        )
    if not doc.get("complete"):
        notes.append(
            "WARNING: this record may be incomplete -- " + "; ".join(doc.get("warnings") or ["parts are missing"]) + "."
        )
    return notes


def suggested_filename(case: dict, ext: str) -> str:
    from case_lookup import clean_title
    slug = re.sub(r"[^a-z0-9]+", "_", clean_title(case.get("title") or "judgment").lower()).strip("_")[:60]
    return f"{slug or 'judgment'}.{ext}"


def build_pdf(case: dict, doc: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    header = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.HexColor("#14213D"))
    meta = ParagraphStyle("m", fontName="Helvetica", fontSize=9.5, leading=13)
    warn = ParagraphStyle("w", fontName="Helvetica-Bold", fontSize=9.5, leading=13, textColor=colors.HexColor("#8B1E1E"))
    disc = ParagraphStyle("d", fontName="Helvetica-Oblique", fontSize=8.8, leading=12)
    body = ParagraphStyle("b", fontName="Times-Roman", fontSize=11, leading=15, spaceAfter=7)

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.grey)
        canvas.drawString(0.8 * inch, 0.5 * inch, "Recourse -- retrieved copy, not independently verified. Confirm against the official source.")
        canvas.drawRightString(A4[0] - 0.8 * inch, 0.5 * inch, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, topMargin=0.8 * inch, bottomMargin=0.85 * inch,
                            leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                            title=_safe(case.get("title") or "Judgment"))
    story = [Paragraph("RECOURSE -- RETRIEVED JUDGMENT", header), Spacer(1, 4),
             HRFlowable(width="100%", thickness=1.2, color=colors.HexColor("#14213D")), Spacer(1, 8)]
    for label, value in _meta_lines(case, doc):
        story.append(Paragraph(f"<b>{label}:</b> {_esc(_safe(value))}", meta))
    story.append(Spacer(1, 8))
    box = Table([[Paragraph(_esc(DISCLAIMER), disc)]], colWidths=[A4[0] - 1.8 * inch])
    box.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#14213D")),
                             ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F1EA")),
                             ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                             ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story.append(box)
    for note in _notices(doc):
        story += [Spacer(1, 6), Paragraph(_esc(note), warn)]
    story += [Spacer(1, 12), HRFlowable(width="100%", thickness=0.5, color=colors.grey), Spacer(1, 10)]
    for para in _paragraphs(_safe(doc.get("text") or "")):
        story.append(Paragraph(_esc(para), body))
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def build_docx(case: dict, doc: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    d = Document()
    d.add_heading("RECOURSE -- RETRIEVED JUDGMENT", level=1)
    for label, value in _meta_lines(case, doc):
        p = d.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)
    p = d.add_paragraph()
    run = p.add_run(DISCLAIMER)
    run.italic = True
    run.font.size = Pt(9)
    for note in _notices(doc):
        p = d.add_paragraph()
        run = p.add_run(note)
        run.bold = True
        run.font.color.rgb = RGBColor(0x8B, 0x1E, 0x1E)
    d.add_paragraph("_" * 60)
    for para in _paragraphs(doc.get("text") or ""):
        d.add_paragraph(para)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
