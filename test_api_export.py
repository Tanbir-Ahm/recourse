"""
test_api_export.py -- GET /api/chats/{id}/export (api/routers/chats.py, answer_export.py): a conversation or one answer as a PDF or a Word file. Free and offline: a throwaway PostgreSQL (the test exits 0 where none exists), a stand-in engine,
the Streamlit service's signed token as the way to be somebody.
Proves: the file is built from the SAVED messages of the person's own conversation (anyone else gets 404, a message id from another chat is 404, a request cannot carry its own content); one answer or the whole conversation, PDF or Word,
with the headers a download needs; every box of the page (currency note, notices, sources) is inside the file; a refusal is included in the conversation file; a question in a script the PDF font lacks gives 422 for PDF and a good Word file; consent
is required; an empty conversation, an unknown format, a deleted chat; the per-person rate limit and the parallel-build cap; an export spends no quota and calls no model.
Run: python -X utf8 test_api_export.py
"""
import io
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
for _v in ("DATABASE_URL", "API_ACCESS_TOKEN", "ACCOUNTS_OPEN", "GOOGLE_CLIENT_ID", "STREAMLIT_API_SECRET", "IP_HASH_SECRET", "CORS_ALLOW_ORIGINS", "TRUSTED_PROXY_HOPS", "API_COOKIE_SECURE", "API_RUN_MIGRATIONS",
           "ACCOUNTS_ALLOW_PLACEHOLDER_LEGAL", "ANON_TRIAL", "ANON_DAILY_BUDGET_INR", "CHAT_MODE", "API_EXPORT_RATE_PER_MINUTE", "API_EXPORT_PARALLEL", "API_RATE_LIMIT_PER_MINUTE"):
    os.environ.pop(_v, None)

import pg_harness

FAILURES = []


def check(cond, msg):
    print(f"[{'PASS' if cond else 'FAIL'}] {msg}")
    if not cond:
        FAILURES.append(msg)


def norm(s):
    return "".join(ch.lower() for ch in str(s or "") if ch.isalnum())


def pdf_text(data):
    import pymupdf
    return "\n".join(p.get_text("text") for p in pymupdf.open(stream=data, filetype="pdf"))


def docx_text(data):
    import docx
    import html
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts += [p.text for p in c.paragraphs]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        body = z.read("word/document.xml").decode("utf-8")
    return "\n".join(parts) + "\n" + "".join(html.unescape(m) for m in re.findall(r"<w:t[^>]*>([^<]*)</w:t>", body))


pg = pg_harness.start()
if pg is None:
    sys.exit(0)
try:
    from fastapi.testclient import TestClient
    from sqlalchemy import func, select

    from api import engine_adapter, ratelimit
    from api.auth import streamlit_token as ST
    from api.db import migrate as MIG
    from api.db import models as M
    from api.db import session as DBS
    from api.main import app
    from api.routers import chats as CH

    SECRET = "s" * 40
    pg.create_db("export")
    os.environ.update({"DATABASE_URL": pg.url("export"), "ACCOUNTS_OPEN": "on", "ACCOUNTS_ALLOW_PLACEHOLDER_LEGAL": "on", "STREAMLIT_API_SECRET": SECRET, "GOOGLE_CLIENT_ID": "client.example", "API_COOKIE_SECURE": "off",
                       "API_RATE_LIMIT_PER_MINUTE": "0"})
    DBS.reset_engine()
    MIG.upgrade(DBS.get_engine())

    class FakeEngine:
        calls = 0

        def answer_question(self, question, inline_domains=frozenset(), audience=None, channel=None):
            import cost_log
            import progress
            FakeEngine.calls += 1
            cost_log.begin(question)
            try:
                progress.emit("scope", category="in_scope")
                last = question.rsplit("Now the person says:", 1)[-1]      # a follow-up reaches the engine folded into the person's earlier questions: only the newest words decide
                if "REFUSE" in last:
                    return {"state": "adjacent_uncovered", "reasoning": "An employment dispute."}
                if "RICH" in last:
                    return {"state": "single_match", "audience": audience, "response_text": "**Your situation**\nYou were arrested.\n\n**What the law says**\nArrest is limited under the BNSS.\n\n**What you can do**\n1. Ask for the grounds.\n2. Call a lawyer.",
                            "matches": [{"case_name": "Arnesh Kumar v. State of Bihar", "citation": "(2014) 8 SCC 273", "paragraph_number": "7_a", "text": "Arrest is not to be made routinely.", "source": "curated", "type": "judgment"}],
                            "register_notes": ["Reported as overruled; not yet confirmed by our reviewer."], "grounding_notes": ["One quotation could not be matched."],
                            "act_notice": {"html": "The <b>BNSS</b> is on <a href=\"https://www.indiacode.nic.in/\">India Code</a>."},
                            "unverified_related_judgments": [{"case_name": "Z v. W", "ik_search_url": "https://indiankanoon.org/search/?formInput=z", "paragraph_number": "2"}]}
                return {"state": "single_match", "response_text": "Plain answer to: " + question[-40:], "matches": [], "audience": audience}
            finally:
                cost_log.finish(style="v1", audience=audience or "")

    engine_adapter.set_engine(FakeEngine())
    engine_adapter.reset_runner(engine_adapter.Runner(4, 8))
    c = TestClient(app)

    def H(sub, email):
        return lambda: {"Authorization": "Bearer " + ST.mint(SECRET, sub, email)}

    A, B, N = H("a-sub", "alice@example.com"), H("b-sub", "bob@example.com"), H("n-sub", "newbie@example.com")

    def make_ready(h):
        for doc in ("privacy", "terms"):
            v = c.get(f"/api/legal/{doc}").json()["version"]
            assert c.post("/api/me/consent", json={"document": doc, "version": v}, headers=h()).status_code == 200
        assert c.post("/api/me/age", json={"confirmed": True}, headers=h()).status_code == 200

    [c.get("/api/me", headers=x()) for x in (A, B, N)]
    make_ready(A)
    make_ready(B)
    keyn = [0]

    def ask(h, chat_id, q):
        keyn[0] += 1
        r = c.post(f"/api/chats/{chat_id}/ask", json={"question": q}, headers={**h(), "Idempotency-Key": f"export-key-{keyn[0]:05d}"})
        assert r.status_code == 200, (r.status_code, r.text[:200])
        return r.json()

    def new_chat(h, audience="layman"):
        r = c.post("/api/chats", json={"audience": audience}, headers=h())
        assert r.status_code == 201, r.text
        return r.json()["id"]

    cid = new_chat(A)
    r1 = ask(A, cid, "RICH: my brother was arrested, what now?")
    r2 = ask(A, cid, "And what about bail?")
    detail = c.get(f"/api/chats/{cid}", headers=A()).json()
    msgs = detail["messages"]
    check(len(msgs) == 4, "a conversation with two questions and two answers is saved")
    a1, a2 = msgs[1], msgs[3]

    def used(h):
        return c.get("/api/me", headers=h()).json()["usage"]


    # ---- the whole conversation
    r = c.get(f"/api/chats/{cid}/export?format=pdf", headers=A())
    check(r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content[:5] == b"%PDF-", f"the conversation as a PDF ({r.status_code}, {r.headers.get('content-type')})")
    check(re.match(r'attachment; filename="recourse-conversation-\d{4}-\d{2}-\d{2}\.pdf"$', r.headers.get("content-disposition", "")) is not None, f"a download file name ({r.headers.get('content-disposition')})")
    check(r.headers.get("cache-control") == "no-store" and r.headers.get("x-content-type-options") == "nosniff", "never cached, never sniffed")
    T = pdf_text(r.content)
    for needle in ("my brother was arrested, what now?", "And what about bail?", "You were arrested.", "Arrest is limited under the BNSS.", "Plain answer to:", "Reported as overruled; not yet confirmed by our reviewer.",
                   "One quotation could not be matched.", "Z v. W", "not independently verified", "Contents", "Question 2", "India Code"):
        check(norm(needle) in norm(T), f"the PDF conversation has: {needle[:60]}")
    rd = c.get(f"/api/chats/{cid}/export?format=docx", headers=A())
    check(rd.status_code == 200 and rd.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" and rd.content[:2] == b"PK", "the conversation as a Word file")
    check(rd.headers["content-disposition"].endswith('.docx"'), "its file name ends in .docx")
    D = docx_text(rd.content)
    for needle in ("my brother was arrested, what now?", "And what about bail?", "Reported as overruled; not yet confirmed by our reviewer.", "Z v. W", "Question 2"):
        check(norm(needle) in norm(D), f"the Word conversation has: {needle[:60]}")

    # ---- one answer
    r = c.get(f"/api/chats/{cid}/export?format=pdf&message={a1['id']}", headers=A())
    one = pdf_text(r.content)
    check(r.status_code == 200 and "recourse-answer-" in r.headers["content-disposition"], "one answer: its own file name")
    check(norm("my brother was arrested") in norm(one) and norm("And what about bail?") not in norm(one) and norm("Reported as overruled") in norm(one), "one answer holds that question and answer and not the other")
    ru = c.get(f"/api/chats/{cid}/export?format=docx&message={msgs[0]['id']}", headers=A())
    check(ru.status_code == 200 and norm("my brother was arrested") in norm(docx_text(ru.content)) and norm("And what about bail?") not in norm(docx_text(ru.content)), "the id of the QUESTION also picks that question and its answer")
    r = c.get(f"/api/chats/{cid}/export?format=pdf&message={a2['id']}", headers=A())
    check(r.status_code == 200 and norm("Plain answer to") in norm(pdf_text(r.content)) and norm("Contents") not in norm(pdf_text(r.content)), "the second answer alone has no contents list")

    # ---- whose it is
    check(c.get(f"/api/chats/{cid}/export", headers=B()).status_code == 404, "another person's conversation is 404 (never 403, never the file)")
    check(c.get(f"/api/chats/{cid}/export").status_code == 401, "not signed in is 401")
    cidb = new_chat(B)
    ask(B, cidb, "Bob's own question")
    check(c.get(f"/api/chats/{cidb}/export?message={a1['id']}", headers=B()).status_code == 404, "an answer id from somebody else's conversation is 404 in one's own")
    check(c.get(f"/api/chats/{cid}/export?message=00000000-0000-0000-0000-000000000000", headers=A()).status_code == 404, "an unknown message id is 404")
    check(c.get(f"/api/chats/{cid}/export?message=not-a-uuid", headers=A()).status_code == 422, "a message id that is not an id is 422")
    check(c.get(f"/api/chats/{cid}/export?format=rtf", headers=A()).status_code == 422, "an unknown format is 422")
    check(c.get(f"/api/chats/{cid}/export", headers=A()).headers["content-type"] == "application/pdf", "PDF is the default format")
    check(c.get(f"/api/chats/{cid}/export?format=pdf", headers=N()).status_code == 403, "a person who has not accepted the notices cannot export (consent_required)")
    empty = new_chat(A)
    check(c.get(f"/api/chats/{empty}/export", headers=A()).status_code == 409 and c.get(f"/api/chats/{empty}/export", headers=A()).json()["detail"] == "nothing_to_export", "an empty conversation has nothing to export (409)")
    check(c.post(f"/api/chats/{cid}/export", headers=A()).status_code == 405, "export is a GET only: nothing can be posted to it")

    # ---- a refusal is part of the conversation
    cid2 = new_chat(A)
    ask(A, cid2, "REFUSE: can I get my job back?")
    ask(A, cid2, "RICH second try")
    r = c.get(f"/api/chats/{cid2}/export?format=pdf", headers=A())
    T2 = pdf_text(r.content)
    check(r.status_code == 200 and norm("can I get my job back?") in norm(T2) and norm("Contents") in norm(T2) and norm("Question 2") in norm(T2), "a conversation with a refused question exports, the refusal included")

    # ---- a script the PDF cannot draw
    cid3 = new_chat(A)
    ask(A, cid3, "मेरे भाई को गिरफ्तार किया गया है")
    rp = c.get(f"/api/chats/{cid3}/export?format=pdf", headers=A())
    check(rp.status_code == 422 and rp.json()["detail"] == "pdf_script_unsupported", f"a Hindi question: 422 pdf_script_unsupported for PDF ({rp.status_code})")
    rw = c.get(f"/api/chats/{cid3}/export?format=docx", headers=A())
    check(rw.status_code == 200 and "मेरे भाई को गिरफ्तार किया गया है" in docx_text(rw.content), "and a Word file that carries the Hindi")

    # ---- an export spends nothing
    calls_after_setup = FakeEngine.calls
    usage_a = used(A)
    for _ in range(3):
        c.get(f"/api/chats/{cid}/export?format=pdf", headers=A())
    check(FakeEngine.calls == calls_after_setup, "an export calls no model")
    check(used(A) == usage_a, "an export uses no daily answer")

    # ---- limits
    ratelimit.reset()
    os.environ["API_EXPORT_RATE_PER_MINUTE"] = "3"
    codes = [c.get(f"/api/chats/{cid}/export?format=docx", headers=A()).status_code for _ in range(5)]
    check(codes == [200, 200, 200, 429, 429], f"the per-person export limit ({codes})")
    r = c.get(f"/api/chats/{cid}/export", headers=A())
    check(r.status_code == 429 and r.json()["detail"] == "rate_limited" and r.headers.get("retry-after"), "over the limit: 429 rate_limited with Retry-After")
    check(c.get(f"/api/chats/{cidb}/export", headers=B()).status_code == 200, "another person's allowance is their own")
    os.environ["API_EXPORT_RATE_PER_MINUTE"] = "0"
    ratelimit.reset()
    os.environ["API_EXPORT_PARALLEL"] = "1"
    slot = CH._export_slot()
    check(slot.acquire(blocking=False), "(setup: take the only build slot)")
    r = c.get(f"/api/chats/{cid}/export", headers=A())
    check(r.status_code == 429 and r.json()["detail"] == "busy", "while the only build slot is taken: 429 busy")
    slot.release()
    check(c.get(f"/api/chats/{cid}/export", headers=A()).status_code == 200, "and a file is built again once the slot is free")
    os.environ.pop("API_EXPORT_PARALLEL", None)
    os.environ.pop("API_EXPORT_RATE_PER_MINUTE", None)

    # ---- deleting a conversation
    check(c.delete(f"/api/chats/{cid}", headers=A()).status_code == 200 and c.get(f"/api/chats/{cid}/export", headers=A()).status_code == 404, "a deleted conversation can no longer be exported")

    # ---- the saved record is what is exported, nothing else: a changed saved answer changes the file
    cid4 = new_chat(A)
    ask(A, cid4, "RICH: before the change")
    d4 = c.get(f"/api/chats/{cid4}", headers=A()).json()["messages"]
    with DBS.session() as s:
        m = s.get(M.Message, d4[1]["id"])
        res = dict(m.result)
        res["response_text"] = "**Your situation**\nA CHANGED SAVED ANSWER."
        m.result = res
        s.commit()
    r = c.get(f"/api/chats/{cid4}/export?format=pdf", headers=A())
    check(norm("A CHANGED SAVED ANSWER") in norm(pdf_text(r.content)), "the file is built from the saved record at the moment of download")
finally:
    pg.stop()

print()
print("ALL PASSED" if not FAILURES else f"{len(FAILURES)} FAILED")
sys.exit(1 if FAILURES else 0)
