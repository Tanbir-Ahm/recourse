"""
test_answer_export.py -- answer_export.py: an answer or a conversation as a PDF or a Word file. Offline and free (no model, no network).
1. the markdown and HTML pieces (inline marks, lists, tables, links, the Act notice);
2. the page's own words: every boilerplate sentence of the export is found in recourse_app.py's source (so a change of wording on the page fails here until the export follows);
3. a rich layman answer and a lawyer's note with the table of authorities: every box of the page is in BOTH files, found again by reading the files (PyMuPDF for the PDF, python-docx and the raw XML for the Word file);
4. a conversation: contents, several answers, a refusal, an answer with no view, an answer with no record;
5. the read-back check catches a missing piece; the plain-PDF fallback; a script the PDF font lacks; control characters; unsafe links;
6. the Word file's structure (element order, relationships) and, where LibreOffice Writer is installed, a real conversion.
Run: python -X utf8 test_answer_export.py
"""
import ast
import datetime
import io
import os
import re
import subprocess
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET

os.environ.setdefault("ANTHROPIC_API_KEY", "dummy")
os.environ.setdefault("VOYAGE_API_KEY", "dummy")
import answer_export as X
import answer_view as V
import authorities_table as AT

FAILURES = []
HERE = os.path.dirname(os.path.abspath(__file__))


def check(cond, msg):
    print(f"[{'PASS' if cond else 'FAIL'}] {msg}")
    if not cond:
        FAILURES.append(msg)


def norm(s):
    return "".join(ch.lower() for ch in str(s or "") if ch.isalnum())


def pdf_text(data):
    import pymupdf
    d = pymupdf.open(stream=data, filetype="pdf")
    return "\n".join(p.get_text("text") for p in d), len(d)


def docx_text(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts += [p.text for p in c.paragraphs]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        body = z.read("word/document.xml").decode("utf-8")
    import html
    return "\n".join(parts) + "\n" + "".join(html.unescape(m) for m in re.findall(r"<w:t[^>]*>([^<]*)</w:t>", body))


# ---------------------------------------------------------------- 1. the small pieces
runs = X.inline("In *Arnesh Kumar v. State of Bihar* (2014) the **Supreme Court** said 5*4 = 20, see [the judgment](https://example.org/j) and `code`.")
check([r["text"] for r in runs if r.get("i")] == ["Arnesh Kumar v. State of Bihar"], "an italic case name is italic")
check([r["text"] for r in runs if r.get("b")] == ["Supreme Court"], "bold is bold")
check("5*4 = 20" in X.plain(runs), "a multiplication sign is not taken for emphasis")
check(any(r.get("href") == "https://example.org/j" for r in runs) and "(https://example.org/j)" in X.plain(runs), "a link keeps its address and prints it after the words")
bad = X.inline("[click](javascript:alert(1)) and [x](data:text/html,y) and [r](/relative)")
check(not any(r.get("href") for r in bad) and "click" in X.plain(bad) and "javascript" not in X.plain(bad), "unsafe link addresses are dropped; the words stay")
check(X.safe_url("https://a.org/x") and not X.safe_url("ftp://a.org") and not X.safe_url('https://a.org/"x') and not X.safe_url("https://a.org/ x") and not X.safe_url(None), "safe_url takes plain http(s) only")

md = ("**What the law says**\n\nSection 35 limits arrest.\nIt continues here.\n\n**Right now:** do this first.\n\n1. Ask for the grounds.\n2. Call a lawyer.\n   Keep the number.\n\n- point one\n* point two\n\n> a quotation\n> goes on\n\n"
      "| Case | Link |\n|---|---|\n| A v B | [open](https://x.org/a) |\n| C \\| D | no link held |\n\n## A heading\nlast")
B = X.md_blocks(md)
kinds = [b["t"] for b in B]
check(kinds == ["h3", "p", "p", "ol", "ul", "quote", "table", "h3", "p"], f"block kinds in order ({kinds})")
check(B[0]["text"] == "What the law says", "a paragraph that is only a bold line becomes a heading")
check(X.plain(B[1]["runs"]) == "Section 35 limits arrest. It continues here.", "a soft line break joins with a space")
check(B[3]["items"][1]["num"] == 2 and "Keep the number." in X.plain(B[3]["items"][1]["runs"]), "numbered items keep their numbers; an indented line continues the item")
check(len(B[4]["items"]) == 2, "both bullet marks make one list")
check(B[5]["t"] == "quote" and "goes on" in X.plain(B[5]["runs"]), "a quotation")
check(B[6]["head"] == ["Case", "Link"] and len(B[6]["rows"]) == 2 and X.plain(B[6]["rows"][1][0]) == "C | D", "a pipe table, an escaped pipe in a cell")
check(any(r.get("href") == "https://x.org/a" for r in B[6]["rows"][0][1]), "a link inside a table cell")
check(X.md_blocks("") == [] and X.md_blocks(None) == [], "empty text gives no blocks")

act = X.html_to_runs('The <b>BNSS</b> is on <a href="https://www.indiacode.nic.in/">India Code</a> &amp; more.<script>alert(1)</script>')
check(X.plain(act).startswith("The BNSS is on India Code (https://www.indiacode.nic.in/) & more."), f"the Act notice keeps its words, bold and link ({X.plain(act)[:80]!r})")
check(any(r.get("b") and r["text"] == "BNSS" for r in act), "bold kept in the Act notice")
act2 = X.html_to_runs('<a href="javascript:alert(1)">go</a> <img src=x onerror=alert(1)>')
check(not any(r.get("href") for r in act2) and "javascript" not in X.plain(act2), "an unsafe link in the Act notice is just words")

# ---------------------------------------------------------------- 2. the page's own words
src = open(os.path.join(HERE, "recourse_app.py"), encoding="utf-8").read()
lits = []
for node in ast.walk(ast.parse(src)):
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        lits.append(node.value)
hay = norm(re.sub(r"&(mdash|middot|nbsp|amp|rsquo|lsquo);", lambda m: {"mdash": "—", "middot": "·", "nbsp": " ", "amp": "&", "rsquo": "'", "lsquo": "'"}[m.group(1)], "\n".join(lits)))
for name in ("CONFLICT", "FALLBACK_LEAD", "H_RELIED", "H_CURRENCY", "H_GROUNDING", "SCI_FOOT", "H_ANCHORS", "ANCHORS_UNCONFIRMED", "ANCHORS_FOOT_1", "ANCHORS_FOOT_2", "ANCHORS_FOOT_3", "H_UNVERIFIED", "UNVERIFIED_FOOT"):
    check(norm(getattr(X, name)) in hay, f"the export's {name} is the page's own wording")
# the wider-search heading and the bucket credit are worded by sci_search_tier (they name the court when a High Court judgment is among the passages); the page calls it, and for Supreme Court passages it says what it said before
import sci_search_tier as _T
check("_SCI.box_title(" in src and "_SCI.credit_for(" in src, "the page words the wider-search heading and the credit through sci_search_tier")
check(_T.box_title([], True) == X.H_SCI_FED and _T.box_title([], False) == X.H_SCI_SHOWN and _T.box_title([{"court": "Supreme Court of India"}], True) == X.H_SCI_FED, "the export's H_SCI_FED / H_SCI_SHOWN are the page's own wording (Supreme Court passages)")
check(_T.credit_for("open_bucket_pdf") == X.OPEN_BUCKET == _T.OPEN_BUCKET_CREDIT, "the export's OPEN_BUCKET is the page's own wording")

# ---------------------------------------------------------------- 3. a rich layman answer
TEXT = ("**Your situation**\nYou were arrested on a *bailable* charge.\n\n**What the law says**\nUnder section 35 of the BNSS an arrest is limited. In *Arnesh Kumar v. State of Bihar* (2014) the Supreme Court said arrest is not routine.\n\n"
        "**What you can do**\n1. Ask for the grounds of arrest.\n2. Call a lawyer.\n\nWhat's unclear: whether you were told the reason, and whether a notice was issued.")
MATCHES = [{"case_name": "Arnesh Kumar v. State of Bihar", "citation": "(2014) 8 SCC 273", "paragraph_number": "7_a", "text": "Arrest is not to be made routinely.", "source": "curated", "type": "judgment"},
           {"section_number": "35", "act": "BNSS", "act_display": "BNSS", "text": "Arrest without warrant, with notice of appearance.", "source": "curated_statute", "type": "statute"}]
RES = {
    "state": "single_match", "audience": "layman", "response_text": TEXT, "matches": MATCHES, "outside_scope_note": "This looks partly like a civil matter.",
    "act_notice": {"html": "The <b>Bharatiya Nagarik Suraksha Sanhita</b> is on <a href=\"https://www.indiacode.nic.in/\">India Code</a>."},
    "register_notes": ["Reported as overruled; not yet confirmed by our reviewer."], "register_possible": ["A possible later case."], "register_not_checked": "Not checked by us.", "grounding_notes": ["One quotation could not be matched to the sources."],
    "cited_judgment_links": [{"case_name": "Arnesh Kumar v. State of Bihar", "url": "https://example.org/a", "source_label": "Indian Kanoon"}, {"case_name": "Bad Link v. State", "url": "javascript:alert(1)", "source_label": "x"}],
    "sci_related_judgments": [{"case_name": "X v. Y", "link": "https://sci.gov.in/x", "date": "2020-01-01", "paragraph": "4", "excerpt": "passage text of x", "passage_note": "a reading note", "register_note": "X was reported as doubted.", "link_kind": "open_bucket_pdf"}],
    "sci_fed": True, "sci_as_of": "2026-10-01",
    "anchor_excerpts": [{"case_name": "P. Ramachandra Rao", "status": "machine_checked", "year": 2002, "link": "https://sci.gov.in/pr", "excerpts": [{"label": "Held", "paragraph": 29, "text": "word for word text", "note": "an anchor note"}]}],
    "unverified_related_judgments": [{"case_name": "Z v. W", "ik_search_url": "https://indiankanoon.org/search/?formInput=z", "paragraph_number": "2", "procedural_disposal": True}],
    "section_note": [{"lead": "Section 17 in full", "texts": ["The text of 17.\nSecond line."], "cites": [{"lead": "Section 13", "texts": ["The text of 13."]}], "loaded_on": "2026-10-17", "cut": True, "source_url": "https://indiacode.nic.in/x"}],
}
VIEW = V.build(RES)
WHEN = datetime.datetime(2026, 10, 23, 9, 30, tzinfo=datetime.timezone.utc)
ENTRY = {"question": "My brother was arrested. What now?", "view": VIEW, "answer_id": "RA-1A2B3C", "created_at": WHEN, "audience": "layman"}
doc = X.build_document([ENTRY], generated_at=WHEN)
check(doc["title"] == "Your question and the answer" and "RA-1A2B3C" in doc["meta"] and "23 Oct 2026, 15:00 IST" in doc["meta"], f"title, answer id and the time in India time ({doc['meta']!r})")
first_note = next(b for b in doc["blocks"] if b["t"] == "note")
check(doc["blocks"].index(first_note) <= 3 and "not legal advice" in X.plain(first_note["runs"]), "the disclaimer comes first")
pdf, info_pdf = X.export(doc, "pdf")
dx, info_dx = X.export(doc, "docx")
check(pdf[:5] == b"%PDF-" and info_pdf["ok"] is True and not info_pdf["fallback"], f"a PDF, proved complete ({info_pdf})")
check(dx[:2] == b"PK" and info_dx["ok"] is True, f"a Word file, proved complete ({info_dx})")
PT, pages = pdf_text(pdf)
DT = docx_text(dx)
MUST = ["My brother was arrested. What now?", "You were arrested on a bailable charge.", "Ask for the grounds of arrest.", "Bharatiya Nagarik Suraksha Sanhita", "https://www.indiacode.nic.in/", "Section 17 in full", "The text of 17.", "Second line.", "Section 13", "The text of 13.", "not yet read by our reviewer",
        "BNSS Section 35", "Arrest without warrant, with notice of appearance.", "Arnesh Kumar v. State of Bihar", "(2014) 8 SCC 273", "paragraph 7", "Arrest is not to be made routinely.",
        X.H_RELIED, "https://example.org/a", "Indian Kanoon", X.H_CURRENCY, "Reported as overruled; not yet confirmed by our reviewer.", "A possible later case.", "Not checked by us.",
        X.H_GROUNDING, "One quotation could not be matched to the sources.", X.H_SCI_FED, "This search does not say whether a decision is still good law", "The collection ends 2026-10-01.", "X v. Y", "paragraph 4", "passage text of x",
        "a reading note", "X was reported as doubted.", "open Supreme Court judgments bucket", X.H_ANCHORS + X.ANCHORS_UNCONFIRMED, "P. Ramachandra Rao", "Held", "word for word text", "an anchor note",
        X.H_UNVERIFIED, "Z v. W", "may be a procedural/bail order, not a full judgment", "Facts to find out or bring to a lawyer", "Whether you were told the reason", "Notes for my lawyer", "Bad Link v. State"]
for needle in MUST:
    check(norm(needle) in norm(PT), f"PDF has: {needle[:70]}")
    check(norm(needle) in norm(DT), f"Word file has: {needle[:70]}")
check("javascript" not in PT and "javascript" not in DT, "the unsafe link's address is nowhere in either file")
check(pages >= 2 and "Page 1 of" in PT, f"the PDF has {pages} pages with 'Page n of N' footers")
with zipfile.ZipFile(io.BytesIO(dx)) as z:
    rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    footer_xml = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/footer"))
check("javascript" not in rels and "https://example.org/a" in rels and "https://sci.gov.in/x" in rels, "the Word file links only to safe addresses")
check("PAGE" in footer_xml and "NUMPAGES" in footer_xml and "information, not legal advice" in footer_xml, "the Word footer carries the notice and live page numbers")
import pymupdf
d_pdf = pymupdf.open(stream=pdf, filetype="pdf")
uris = [l.get("uri") for p in d_pdf for l in p.get_links() if l.get("uri")]
check("https://example.org/a" in uris and not any(u.lower().startswith("javascript") for u in uris), f"PDF links are clickable and safe ({len(uris)} links)")
check(d_pdf.metadata.get("title", "").startswith("Recourse") and d_pdf.metadata.get("author") == "Recourse", "PDF title and author are set")
import docx as _docx
cp = _docx.Document(io.BytesIO(dx)).core_properties
check(cp.author == "Recourse" and cp.title.startswith("Recourse"), "Word title and author are set")

check(VIEW["outside_scope_note"] is None, "with an Act notice the page hides the scope note (the page's rule), so the export does too")
scope_view = V.build(dict(RES, act_notice=None))
sdoc = X.build_document([dict(ENTRY, view=scope_view)], generated_at=WHEN)
sP, sD = X.export(sdoc, "pdf")[0], X.export(sdoc, "docx")[0]
check(norm("This looks partly like a civil matter.") in norm(pdf_text(sP)[0]) and norm("This looks partly like a civil matter.") in norm(docx_text(sD)), "without an Act notice the outside-scope note is in both files")
cdoc_conf = X.build_document([dict(ENTRY, view=V.build(dict(RES, state="conflicting_matches")))], generated_at=WHEN)
check(norm(X.CONFLICT) in norm(X.all_text(cdoc_conf)) or X.CONFLICT in "".join(X.plain(b.get("runs", [])) for b in cdoc_conf["blocks"] if b["t"] == "p"), "a conflicting answer carries the page's 'more than one provision' note")

# ---------------------------------------------------------------- the lawyer's note with the table of authorities
status = lambda name, year, cit: {"source": "register", "note": "Reported as overruled; not yet confirmed by our reviewer."} if name.startswith("Arnesh") else {"source": "not_checked", "note": None}
AUTH = AT.build(MATCHES, ["Arnesh Kumar v. State of Bihar"], RES["cited_judgment_links"][:1], RES["sci_related_judgments"], status, anchors=RES["anchor_excerpts"])
LRES = dict(RES, audience="lawyer", authorities=AUTH, outside_scope_note=None, act_notice=None)
LVIEW = V.build(LRES)
check(LVIEW["authorities"] and LVIEW["authorities"]["tables"], "the lawyer's view carries the tables of authorities")
ldoc = X.build_document([dict(ENTRY, view=LVIEW, audience="lawyer")], generated_at=WHEN)
check(ldoc["title"] == "Research note" and X.DISCLAIMER_LAWYER in X.plain(next(b for b in ldoc["blocks"] if b["t"] == "note")["runs"]), "the lawyer's file is a Research note with the lawyer's disclaimer")
check(not any(b["t"] == "lines" for b in ldoc["blocks"]) and "Notes for my lawyer" not in X.all_text(ldoc), "no 'Notes for my lawyer' lines in a lawyer's file")
check(sum(1 for b in ldoc["blocks"] if b["t"] == "table") == len(LVIEW["authorities"]["tables"]), "one table per table of authorities")
lpdf, lip = X.export(ldoc, "pdf")
ldx, lid = X.export(ldoc, "docx")
check(lip["ok"] is True and lid["ok"] is True, f"both lawyer files proved complete ({lip}, {lid})")
LP, LD = pdf_text(lpdf)[0], docx_text(ldx)
for needle in ["Where it stands on our lists", "Authorities (reviewed library)", "Further reading from a wider search (not independently verified)", "Arnesh Kumar v. State of Bihar (2014)", "(2014) 8 SCC 273", "Reported as overruled; not yet confirmed by our reviewer.",
               "Please ask a lawyer before relying on any case here.", "Facts to find out"]:
    ok = norm(needle) in norm(LP) and norm(needle) in norm(LD)
    if needle == "Facts to find out":
        ok = norm(needle) not in norm(LP) and norm(needle) not in norm(LD)
        check(ok, "the layman's checklist is not in a lawyer's file")
    else:
        check(ok, f"lawyer's file (both formats) has: {needle[:70]}")
import docx
tbls = docx.Document(io.BytesIO(ldx)).tables
check(len(tbls) >= 2 and tbls[0].rows[0].cells[0].text == "Case", "the Word file holds real tables with a header row")

# ---------------------------------------------------------------- 4. a conversation
NOTICE_VIEW = V.build({"state": "adjacent_uncovered", "reasoning": "An employment dispute.", "register_notes": ["Case A was reported as overruled."]})
check(NOTICE_VIEW and NOTICE_VIEW["kind"] == "notice", "a refusal has a notice view")
entries = [
    dict(ENTRY, answer_id="RA-000001"),
    {"question": "Can I get a job back after dismissal?", "view": NOTICE_VIEW, "answer_id": "RA-000002", "created_at": WHEN, "audience": "layman"},
    {"question": "And what about bail?", "view": None, "plain": "**Short answer**\n\nBail depends on the offence.", "answer_id": "RA-000003", "created_at": WHEN, "audience": "layman"},
    {"question": "A question with no saved answer", "view": None, "plain": None, "answer_id": "RA-000004", "created_at": WHEN, "audience": "layman"},
]
cdoc = X.build_document(entries, kind="conversation", title_chat="My brother's arrest", generated_at=WHEN)
check(cdoc["title"] == "Your conversation with Recourse" and "4 questions" in cdoc["meta"], "a conversation file")
cpdf, cpi = X.export(cdoc, "pdf")
cdx, cdi = X.export(cdoc, "docx")
check(cpi["ok"] is True and cdi["ok"] is True, f"both conversation files proved complete ({cpi}, {cdi})")
CP, CD = pdf_text(cpdf)[0], docx_text(cdx)
for needle in ["Contents", "Question 1", "Question 4", "RA-000003", "Conversation: My brother's arrest", "Can I get a job back after dismissal?", "Case A was reported as overruled.", X.H_CURRENCY, "Bail depends on the offence.",
               "No answer was saved for this question.", "Short answer"]:
    check(norm(needle) in norm(CP) and norm(needle) in norm(CD), f"conversation file (both formats) has: {needle[:60]}")
check(len(re.findall(r"Notes for my lawyer", CP)) == 1, "'Notes for my lawyer' appears once, at the end")
many = X.build_document([dict(ENTRY, answer_id=f"RA-{i:06d}", question=f"Question number {i}") for i in range(20)], kind="conversation", generated_at=WHEN)
check(many["count"] == X.MAX_ENTRIES and "RA-000019" in X.all_text(many) and "RA-000007" not in X.all_text(many), "more than the cap keeps the newest answers")
try:
    X.build_document([])
    check(False, "no entries is refused")
except X.ExportError:
    check(True, "no entries is refused")

# ---------------------------------------------------------------- 5. the read-back check and the fallbacks
other = X.build_document([dict(ENTRY, question="A completely different question about tenancy", answer_id="RA-999999")], generated_at=WHEN)
v1, v2 = X.verify_pdf(pdf, other), X.verify_docx(dx, other)
check(v1["ok"] is False and v1["missing"], f"the PDF check finds missing text ({len(v1['missing'])} pieces)")
check(v2["ok"] is False and v2["missing"], f"the Word check finds missing text ({len(v2['missing'])} pieces)")
check(X.verify_pdf(b"not a pdf", doc)["ok"] is None and X.verify_docx(b"not a docx", doc)["ok"] is None, "an unreadable file gives 'cannot check', not a pass")

real_verify = X.verify_pdf
calls = []
X.verify_pdf = lambda data, d: (calls.append(1), {"ok": False, "checked": 1, "missing": ["x"]} if len(calls) == 1 else real_verify(data, d))[1]
try:
    data, info = X.export(doc, "pdf")
    check(info["fallback"] is True and info["ok"] is True and norm("You were arrested on a bailable charge") in norm(pdf_text(data)[0]), "a failed check rebuilds the PDF in the plain layout")
    X.verify_pdf = lambda data, d: {"ok": False, "checked": 1, "missing": ["x"]}
    try:
        X.export(doc, "pdf")
        check(False, "two failed checks must refuse")
    except X.ExportError:
        check(True, "a PDF that fails its check twice is refused, never sent")
finally:
    X.verify_pdf = real_verify
real_vd = X.verify_docx
X.verify_docx = lambda data, d: {"ok": False, "checked": 1, "missing": ["x"]}
try:
    X.export(doc, "docx")
    check(False, "a Word file that fails its check must be refused")
except X.ExportError:
    check(True, "a Word file that fails its check is refused")
finally:
    X.verify_docx = real_vd
try:
    X.export(doc, "rtf")
    check(False, "unknown format refused")
except X.ExportError:
    check(True, "unknown format refused")

hindi = X.build_document([dict(ENTRY, question="मेरे भाई को गिरफ्तार किया गया है", view=VIEW)], generated_at=WHEN)
check(X.uncovered_characters(hindi), "Devanagari is a script the PDF font lacks")
try:
    X.export(hindi, "pdf")
    check(False, "a PDF with Devanagari must be refused")
except X.ScriptNotSupported:
    check(True, "a PDF with letters its font cannot draw is refused, not sent as empty boxes")
hd, hdi = X.export(hindi, "docx")
check(hdi["ok"] is True and "मेरे भाई को गिरफ्तार किया गया है" in docx_text(hd), "the Word file carries the Hindi question")
check(not X.uncovered_characters(doc), "an English answer with curly quotes, dashes and the rupee sign is fully covered")
rupee = X.build_document([dict(ENTRY, view=V.build(dict(RES, response_text="The fine is ₹5,000 — “quoted” → next ☐.")))], generated_at=WHEN)
check(not X.uncovered_characters(rupee) and X.export(rupee, "pdf")[1]["ok"] is True, "rupee sign, dashes, curly quotes, arrows and check boxes print")

ctrl = dict(RES, sci_related_judgments=[{"case_name": "Scan v. Text", "excerpt": "page one\x0cpage two \x00 bad \x08 chars", "link": "https://sci.gov.in/s"}])
cview = V.build(ctrl)
cdoc2 = X.build_document([dict(ENTRY, view=cview)], generated_at=WHEN)
check("\x00" not in X.all_text(cdoc2) and "\x0c" not in X.all_text(cdoc2) and "page one\npage two" in X.all_text(cdoc2), "control characters are removed and a form feed becomes a line break")
check(X.export(cdoc2, "pdf")[1]["ok"] is True and X.export(cdoc2, "docx")[1]["ok"] is True, "a scanned excerpt with control characters builds in both formats")
odd = X.build_document([dict(ENTRY, question="<script>alert(1)</script> & <b>x</b> \"quoted\"", view=V.build(dict(RES, response_text="Text with <tags> & ampersands & 'quotes' and a very long url https://example.org/" + "a" * 400)))], generated_at=WHEN)
oP, oD = X.export(odd, "pdf"), X.export(odd, "docx")
check(oP[1]["ok"] is True and oD[1]["ok"] is True and "<script>alert(1)</script>" in pdf_text(oP[0])[0] and "<script>alert(1)</script>" in docx_text(oD[0]), "markup characters in the text are printed as text, never interpreted")

# ---------------------------------------------------------------- size and speed
big_text = "Section text line. " * 800
bigres = dict(RES, matches=MATCHES + [{"section_number": "17", "act": "PCACT", "act_display": "PC Act", "text": big_text, "source": "act_shelf", "retrieval": "named", "shelf": {"loaded_on": "2026-10-17"}}], response_text=TEXT + "\n\nSection 17 of the PC Act matters.")
bdoc = X.build_document([dict(ENTRY, view=V.build(bigres))] * 3, kind="conversation", generated_at=WHEN)
t0 = time.time()
bp, bpi = X.export(bdoc, "pdf")
t1 = time.time()
bd, bdi = X.export(bdoc, "docx")
t2 = time.time()
check(bpi["ok"] is True and bdi["ok"] is True, f"three answers with a 15,000 character section each: both proved complete ({bpi['checked']} pieces)")
check(t1 - t0 < 15 and t2 - t1 < 15, f"built in a reasonable time (PDF {t1 - t0:.1f} s, Word {t2 - t1:.1f} s)")

# ---------------------------------------------------------------- 6. the Word file's structure
def ordered_ok(xml_bytes):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    W = "{%s}" % ns["w"]
    root = ET.fromstring(xml_bytes)
    order_p = [t.replace("w:", W) for t in X._PPR_ORDER]
    order_r = [W + t for t in ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
                               "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath")]
    order_tc = [W + t for t in ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark")]
    bad = []
    for tag, order in (("pPr", order_p), ("rPr", order_r), ("tcPr", order_tc)):
        for el in root.iter(W + tag):
            idx = [order.index(c.tag) if c.tag in order else -1 for c in el]
            if -1 in idx:
                bad.append((tag, [c.tag.split("}")[1] for c in el if c.tag not in order]))
            elif idx != sorted(idx):
                bad.append((tag, [c.tag.split("}")[1] for c in el]))
    return bad


with zipfile.ZipFile(io.BytesIO(dx)) as z:
    check(z.testzip() is None, "the Word file is a sound zip")
    names = z.namelist()
    check("[Content_Types].xml" in names and "word/document.xml" in names, "the Word file has its required parts")
    bad_order = ordered_ok(z.read("word/document.xml"))
    check(not bad_order, f"paragraph, run and cell properties are in the order Word requires ({bad_order[:3]})")
    bad_order_big = ordered_ok(zipfile.ZipFile(io.BytesIO(ldx)).read("word/document.xml"))
    check(not bad_order_big, f"the same for the lawyer's file with tables ({bad_order_big[:3]})")

soffice = None
for cand in ("soffice", "libreoffice"):
    try:
        if subprocess.run([cand, "--version"], capture_output=True, timeout=60).returncode == 0:
            soffice = cand
            break
    except Exception:  # noqa: BLE001
        pass
if soffice:
    with tempfile.TemporaryDirectory() as tmp:
        for name, data in (("answer.docx", dx), ("note.docx", ldx), ("chat.docx", cdx)):
            with open(os.path.join(tmp, name), "wb") as f:
                f.write(data)
        env = dict(os.environ, HOME=tmp)
        r = subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "pdf", "--outdir", tmp] + [os.path.join(tmp, n) for n in ("answer.docx", "note.docx", "chat.docx")],
                           capture_output=True, timeout=240, env=env)
        got = [n for n in ("answer.pdf", "note.pdf", "chat.pdf") if os.path.exists(os.path.join(tmp, n))]
        if len(got) == 3:
            check(True, "LibreOffice opens and converts all three Word files")
            t_answer = pdf_text(open(os.path.join(tmp, "answer.pdf"), "rb").read())[0]
            check(all(norm(k) in norm(t_answer) for k in ("You were arrested on a bailable charge.", X.H_CURRENCY, X.H_UNVERIFIED, "Page 1 of")), "the converted Word file shows the answer, the warnings and live page numbers")
            t_chat = pdf_text(open(os.path.join(tmp, "chat.pdf"), "rb").read())[0]
            check(all(norm(k) in norm(t_chat) for k in ("Contents", "Question 4", "Bail depends on the offence.")), "the converted conversation shows all its questions")
        else:
            print("[SKIP] LibreOffice could not convert the Word files here (Writer not installed?):", (r.stderr or b"")[-200:])
else:
    print("[SKIP] LibreOffice is not installed: no rendering check of the Word files")

print()
print("ALL PASSED" if not FAILURES else f"{len(FAILURES)} FAILED")
sys.exit(1 if FAILURES else 0)
